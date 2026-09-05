"""بناء ملف تسليم دليل الاستخدام (PDF + DOCX) من manual_source.html."""
import html as html_lib
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from playwright.sync_api import sync_playwright

BASE = Path(__file__).resolve().parent
SRC = BASE / "deliverables" / "manual_source.html"
OUT_PDF = BASE / "deliverables" / "دليل_استخدام_النظام.pdf"
OUT_DOCX = BASE / "deliverables" / "دليل_استخدام_النظام.docx"

PRIMARY = RGBColor(0x1A, 0x56, 0xDB)


def build_pdf():
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(SRC.resolve().as_uri())
        page.wait_for_timeout(200)
        page.pdf(
            path=str(OUT_PDF),
            format="A4",
            print_background=True,
            margin={"top": "12mm", "bottom": "12mm", "left": "10mm", "right": "10mm"},
            display_header_footer=True,
            header_template='<div style="width:100%;text-align:center;font-size:8pt;color:#94a3b8;padding:0 10mm;">نظام فواتير المياه — دليل الاستخدام</div>',
            footer_template='<div style="width:100%;text-align:center;font-size:8pt;color:#94a3b8;padding:0 10mm;"><span class="pageNumber"></span> / <span class="totalPages"></span></div>',
        )
        browser.close()
    print("PDF:", OUT_PDF, f"({OUT_PDF.stat().st_size} bytes)")


def _txt(raw: str) -> str:
    return html_lib.unescape(raw).replace("\xa0", " ").strip()


class ManualParser(HTMLParser):
    """يحوّل manual_source.html إلى عناصر (عناوين/فقرات/قوائم/جداول/ملاحظات)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.elements = []  # dicts: kind, text, extra
        self._skip_depth = 0
        self._buf = []
        self._in_sec_title = 0
        self._in_note = 0
        self._li_kind = None  # 'ol' | 'ul'
        self._li_ol_index = 0
        self._li_buf = []
        self._in_table = 0
        self._row = None  # None أو قائمة خلايا
        self._cell_buf = []
        self._in_h4 = 0
        self._in_p = 0
        self._after_li_end = False
        self._in_sel_number = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script"):
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "div" and attrs and any(k == "class" and v == "sec-title" for k, v in attrs):
            self._in_sec_title += 1
            self._buf = []
        elif tag == "div" and attrs and any(k == "class" and v == "note" for k, v in attrs):
            self._in_note += 1
            self._buf = []
        elif tag == "span" and attrs and any(k == "class" and v == "num" for k, v in attrs):
            self._in_sel_number += 1
        elif tag in ("h1", "h2") and self._skip_depth == 0:
            self._buf = []
        elif tag == "h4":
            self._in_h4 += 1
            self._buf = []
        elif tag == "p" and not self._in_sec_title:
            self._in_p += 1
            self._buf = []
        elif tag in ("ol", "ul"):
            self._li_kind = tag
            self._li_ol_index = 0
        elif tag == "li":
            self._li_buf = []
        elif tag == "table":
            self._in_table += 1
            self._row = None
        elif tag == "tr" and self._in_table:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell_buf = []

    def handle_endtag(self, tag):
        if self._skip_depth:
            if tag in ("style", "script"):
                self._skip_depth -= 1
            return
        if tag == "span" and self._in_sel_number:
            self._in_sel_number -= 1
        elif tag == "div" and self._in_sec_title:
            self._in_sec_title -= 1
            self.elements.append({"kind": "title", "text": _txt("".join(self._buf))})
            self._buf = []
        elif tag == "div" and self._in_note:
            self._in_note -= 1
            self.elements.append({"kind": "note", "text": _txt("".join(self._buf))})
            self._buf = []
        elif tag == "h4" and self._in_h4:
            self._in_h4 -= 1
            self.elements.append({"kind": "h4", "text": _txt("".join(self._buf))})
            self._buf = []
        elif tag == "p" and self._in_p:
            self._in_p -= 1
            self.elements.append({"kind": "p", "text": _txt("".join(self._buf))})
            self._buf = []
        elif tag == "li":
            if self._li_kind == "ol":
                self._li_ol_index += 1
                prefix = f"{self._li_ol_index}. "
            else:
                prefix = "• "
            self.elements.append({"kind": "li", "text": prefix + _txt("".join(self._li_buf))})
            self._li_buf = []
        elif tag in ("ol", "ul"):
            self._li_kind = None
        elif tag == "td" or tag == "th":
            if self._row is not None:
                self._row.append(_txt("".join(self._cell_buf)))
            self._cell_buf = []
        elif tag == "tr":
            if self._row is not None and self._in_table:
                self.elements.append({"kind": "row", "cells": self._row})
            self._row = None
        elif tag == "table":
            self._in_table -= 1

    def handle_data(self, data):
        if self._skip_depth:
            return
        if self._in_table and self._row is not None and self._cell_buf is not None:
            self._cell_buf.append(data)
            return
        if self._row is None and self._in_table and self._in_sec_title == 0:
            return
        if self._in_sec_title:
            self._buf.append(data)
        elif self._in_note:
            self._buf.append(data)
        elif self._in_h4:
            self._buf.append(data)
        elif self._in_p:
            self._buf.append(data)
        elif self._li_kind:
            self._li_buf.append(data)


def build_docx():
    doc = Document()
    # الغلاف
    for text, size in [
        ("\n\n\nدليل استخدام نظام فواتير المياه", 26),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = PRIMARY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("دليل عملي خطوة بخطوة لإدارة المشتركين والقراءات والفواتير والتحصيل والتقارير")
    run.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\nبرمجة وتصميم: يمن كود للتقنيات الذكية — 2026")
    run.font.size = Pt(11)
    doc.add_page_break()

    body = SRC.read_text(encoding="utf-8").split("<body>", 1)[1].split("</body>", 1)[0]
    parser = ManualParser()
    parser.feed(body)
    parser.close()

    table_rows = []
    for el in parser.elements:
        kind = el["kind"]
        if kind == "title":
            h = doc.add_paragraph()
            run = h.add_run(el["text"])
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = PRIMARY
        elif kind == "h4":
            h = doc.add_paragraph()
            run = h.add_run(el["text"])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = PRIMARY
        elif kind == "p":
            p = doc.add_paragraph()
            run = p.add_run(el["text"])
            run.font.size = Pt(11.5)
        elif kind == "note":
            p = doc.add_paragraph()
            run = p.add_run("ملاحظة: " + el["text"])
            run.font.size = Pt(11)
            run.italic = True
        elif kind == "li":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(10)
            run = p.add_run(el["text"])
            run.font.size = Pt(11.5)
        elif kind == "row":
            table_rows.append(el["cells"])
        elif kind == "spacer":
            doc.add_paragraph()

    if table_rows:
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=0, cols=cols)
        table.style = "Table Grid"
        first = True
        for row in table_rows:
            cells_row = table.add_row().cells
            for i, val in enumerate(row):
                cell = cells_row[i]
                cell.text = val
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10.5)
                        run.bold = first
            first = False

    doc.save(str(OUT_DOCX))
    print("DOCX:", OUT_DOCX, f"({OUT_DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print("تم بناء ملفي التسليم بنجاح.")